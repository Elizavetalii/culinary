from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "data_dictionary.md"
OUTPUT = ROOT / "docs" / "data_dictionary.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=8):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:type"), "dxa")
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))


def parse_markdown_table(path):
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            continue
        parts = [part.strip() for part in line.strip().strip("|").split("|")]
        if not parts or all(part.startswith("---") for part in parts):
            continue
        rows.append(parts)
    return rows


def build_docx():
    rows = parse_markdown_table(SOURCE)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Таблица 1 - Словарь данных информационной системы ArtCulinaryCRM")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(14)

    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(8)
    note_run = note.add_run(
        "Обозначения ключей: PK - первичный ключ, FK - внешний ключ, UK - уникальный ключ."
    )
    note_run.font.name = "Arial"
    note_run.font.size = Pt(9)

    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    widths = [1.6, 4.0, 3.2, 4.7, 13.6]
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(["Ключ", "Поле", "Тип данных", "Обязательность заполнения", "Описание"]):
        set_cell_text(header.cells[idx], text, bold=True, size=8)
        set_cell_shading(header.cells[idx], "D9EAF7")
        set_cell_width(header.cells[idx], widths[idx])

    for source_row in rows[1:]:
        if len(source_row) != 5:
            continue
        row = table.add_row()
        is_table_caption = source_row[1].startswith("**Таблица")
        for idx, text in enumerate(source_row):
            clean = text.replace("**", "").replace("`", "")
            set_cell_text(row.cells[idx], clean, bold=is_table_caption, size=7.5)
            set_cell_width(row.cells[idx], widths[idx])
            if is_table_caption:
                set_cell_shading(row.cells[idx], "F2F4F7")
            elif idx == 0 and clean:
                set_cell_shading(row.cells[idx], "F8FBFD")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
