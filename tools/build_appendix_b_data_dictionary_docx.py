from pathlib import Path

import os

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "factory_crm.settings")

import django
from django.db import models
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


django.setup()

from django.apps import apps  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Приложение Б - Словарь данных.docx"
APP_LABELS = {"admin_panel", "communications", "crm", "reports"}

MODEL_DESCRIPTIONS = {
    "admin_panel_backup": "Резервная копия",
    "admin_panel_backupschedule": "Расписание резервного копирования",
    "reports_report": "Отчёт",
}

FIELD_DESCRIPTIONS = {
    "id": "Идентификатор записи",
    "password": "Хэш пароля",
    "username": "Логин пользователя",
    "first_name": "Имя",
    "last_name": "Фамилия",
    "email": "Email",
    "full_name": "ФИО",
    "phone": "Телефон",
    "last_login": "Дата последнего входа",
    "is_superuser": "Признак суперпользователя",
    "is_staff": "Признак доступа к административной панели",
    "is_active": "Признак активной учётной записи",
    "date_joined": "Дата регистрации",
    "object_id": "Идентификатор связанного объекта",
    "assigned_at": "Дата назначения",
}


def field_type(field):
    if isinstance(field, models.BigAutoField):
        return "Bigint"
    if isinstance(field, models.AutoField):
        return "Integer"
    if isinstance(field, (models.ForeignKey, models.OneToOneField, models.BigIntegerField)):
        return "Bigint"
    if isinstance(field, models.PositiveSmallIntegerField):
        return "Smallint"
    if isinstance(field, models.PositiveIntegerField):
        return "Integer"
    if isinstance(field, models.IntegerField):
        return "Integer"
    if isinstance(field, models.BooleanField):
        return "Boolean"
    if isinstance(field, models.EmailField):
        return f"Varchar({field.max_length or 254})"
    if isinstance(field, models.URLField):
        return f"Varchar({field.max_length or 200})"
    if isinstance(field, models.FileField):
        return f"Varchar({field.max_length or 100})"
    if isinstance(field, models.CharField):
        return f"Varchar({field.max_length})"
    if isinstance(field, models.TextField):
        return "Text"
    if isinstance(field, models.DateTimeField):
        return "Timestamp"
    if isinstance(field, models.DateField):
        return "Date"
    if isinstance(field, models.TimeField):
        return "Time"
    if isinstance(field, models.DecimalField):
        return f"Decimal({field.max_digits},{field.decimal_places})"
    if isinstance(field, models.JSONField):
        return "JSON"
    return field.get_internal_type()


def field_key(field):
    keys = []
    if field.primary_key:
        keys.append("PK")
    if isinstance(field, (models.ForeignKey, models.OneToOneField)):
        keys.append("FK")
    if field.unique and not field.primary_key:
        keys.append("UK")
    return ", ".join(keys)


def required(field):
    value = "Null" if field.null else "Not null"
    if field.unique and not field.primary_key:
        value += ", unique"
    return value


def field_name(field):
    if isinstance(field, (models.ForeignKey, models.OneToOneField)):
        return field.attname
    return field.name


def description(field):
    if isinstance(field, (models.ForeignKey, models.OneToOneField)):
        target = field.remote_field.model._meta.db_table
        text = f"Ссылка на {target}"
    else:
        text = FIELD_DESCRIPTIONS.get(field.name, str(field.verbose_name))
        if text:
            text = text[0].upper() + text[1:]
    if field.choices:
        values = " / ".join(str(item[0]) for item in field.choices)
        text = f"{text} ({values})"
    if getattr(field, "auto_now_add", False):
        text = f"{text}, заполняется автоматически при создании"
    elif getattr(field, "auto_now", False):
        text = f"{text}, обновляется автоматически"
    return text


def table_rows():
    result = []
    models_list = [
        model
        for model in apps.get_models()
        if model._meta.app_label in APP_LABELS and not model._meta.auto_created
    ]
    models_list.sort(key=lambda model: (model._meta.app_label, model._meta.db_table))

    for model in models_list:
        model_description = MODEL_DESCRIPTIONS.get(model._meta.db_table, str(model._meta.verbose_name))
        result.append([f"Таблица {model._meta.db_table}", "", "", "", model_description])
        for field in model._meta.fields:
            result.append(
                [
                    field_key(field),
                    field_name(field),
                    field_type(field),
                    required(field),
                    description(field),
                ]
            )
        for group in model._meta.unique_together or []:
            if isinstance(group, (tuple, list)) and group:
                columns = []
                for name in group:
                    field = model._meta.get_field(name)
                    columns.append(field_name(field))
                result.append(
                    [
                        "UK",
                        f"({', '.join(columns)})",
                        "Composite",
                        "Not null, unique",
                        f"Уникальное сочетание полей таблицы {model._meta.db_table}",
                    ]
                )
    return result


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_text(cell, text):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    run.font.size = Pt(12)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build():
    document = Document()
    section = document.sections[0]
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Таблица 1 - Словарь данных")
    title_run.font.size = Pt(12)

    table = document.add_table(rows=2, cols=5)
    table.style = "Table Grid"
    table.autofit = True

    headers = ["Ключ", "Поле", "Тип данных", "Обязательность заполнения", "Описание"]
    numbers = ["1", "2", "3", "4", "5"]
    for index, value in enumerate(headers):
        set_text(table.rows[0].cells[index], value)
    for index, value in enumerate(numbers):
        set_text(table.rows[1].cells[index], value)
    set_repeat_header(table.rows[0])
    set_repeat_header(table.rows[1])

    for row_values in table_rows():
        row = table.add_row()
        is_caption = row_values[0].startswith("Таблица ")
        if is_caption:
            merged = row.cells[0].merge(row.cells[4])
            set_text(merged, row_values[0])
        else:
            for index, value in enumerate(row_values):
                set_text(row.cells[index], value)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
